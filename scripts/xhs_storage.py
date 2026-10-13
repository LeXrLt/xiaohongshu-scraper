"""PostgreSQL 存储与 Markdown / CSV / JSON / XLSX 导出。"""

from __future__ import annotations

import csv
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from database import Connection, connect as _connect
from psycopg2.extras import DictRow
from xhs_config import DATA_DIR, ROOT, MEDIA_DIR, OUTPUT_DIR, note_media_dir, sanitize_filename


def _atomic_write(path: Path, content: str, encoding: str = "utf-8") -> None:
    """原子写入文件：先写临时文件再 rename，防止崩溃时文件损坏。"""
    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        tmp.write_text(content, encoding=encoding)
        tmp.replace(path)
    except Exception:
        # replace 失败时清理临时文件
        try:
            tmp.unlink()
        except OSError:
            pass
        raise


def connect() -> Connection:
    return _connect(initialize=True)


def _json_value(value: Any, default: Any) -> Any:
    if isinstance(value, str):
        return json.loads(value) if value else default
    return default if value is None else value


def upsert_note(conn: Connection, note: dict[str, Any]) -> None:
    raw = _json_value(note.get("raw"), {})
    data = {key: note.get(key, "") or "" for key in (
        "note_id", "user_id", "title", "description", "type", "ip_location",
        "published_at", "xsec_token", "xsec_source", "cover_url", "video_transcript",
        "video_ocr_text", "video_summary", "image_ocr_text", "image_summary",
        "image_mermaid", "content_hash",
    )}
    data.update({key: int(note.get(key, 0) or 0) for key in (
        "liked_count", "collected_count", "comment_count", "share_count", "video_duration",
    )})
    data.update(
        topics=json.dumps(_json_value(note.get("topics"), []), ensure_ascii=False),
        raw_data=json.dumps(raw, ensure_ascii=False),
        content=raw.get("desc") or note.get("content") or note.get("description") or "",
        note_url=note.get("note_url") or f"https://www.xiaohongshu.com/explore/{note['note_id']}",
        image_urls=json.dumps(_extract_images(raw), ensure_ascii=False),
        video_url=note.get("video_url") or _extract_video_url(raw) or "",
    )
    preserved = {
        "xsec_token", "xsec_source", "cover_url", "video_url", "video_transcript",
        "video_ocr_text", "video_summary", "image_ocr_text", "image_summary",
        "image_mermaid", "content", "note_url",
    }
    updates = []
    for key in data:
        if key == "note_id":
            continue
        value = f"EXCLUDED.{key}"
        if key in preserved:
            value = f"COALESCE(NULLIF({value}, ''), xhs_notes.{key})"
        elif key == "video_duration":
            value = "CASE WHEN EXCLUDED.video_duration > 0 THEN EXCLUDED.video_duration ELSE xhs_notes.video_duration END"
        elif key == "raw_data":
            value = "CASE WHEN length(EXCLUDED.raw_data::text) > length(COALESCE(xhs_notes.raw_data, '{}'::jsonb)::text) THEN EXCLUDED.raw_data ELSE xhs_notes.raw_data END"
        elif key == "image_urls":
            value = "CASE WHEN EXCLUDED.image_urls <> '[]'::jsonb THEN EXCLUDED.image_urls ELSE xhs_notes.image_urls END"
        updates.append(f"{key}={value}")
    updates.extend([
        "crawled_at=CASE WHEN EXCLUDED.content_hash <> '' AND EXCLUDED.content_hash = xhs_notes.content_hash THEN xhs_notes.crawled_at ELSE CURRENT_TIMESTAMP END",
        "updated_at=CURRENT_TIMESTAMP",
    ])
    with conn:
        conn.execute(
            f"INSERT INTO xhs_notes ({','.join(data)}) VALUES ({','.join(['%s'] * len(data))}) "
            f"ON CONFLICT (note_id) DO UPDATE SET {','.join(updates)}",
            list(data.values()),
        )
        update_media_state(conn, note["note_id"])


def update_media_state(conn: Connection, note_id: str) -> None:
    directory = _find_media_dir(note_id, conn)
    has_media = directory.is_dir() and any(directory.iterdir())
    media_path = ""
    if has_media:
        base = DATA_DIR if directory.is_relative_to(DATA_DIR) else ROOT
        media_path = str(directory.relative_to(base)).replace("\\", "/")
    conn.execute(
        "UPDATE xhs_notes SET media_path=%s, has_local_media=%s WHERE note_id=%s",
        (media_path, has_media, note_id),
    )


def update_video_analysis(
    conn: Connection,
    note_id: str,
    transcript: str = "",
    ocr_text: str = "",
    summary: str = "",
) -> None:
    """更新视频分析结果（转录 / OCR / 摘要），并刷新 crawled_at，不影响其他字段。"""
    with conn:
        conn.execute(
            "UPDATE xhs_notes SET video_transcript=%s, video_ocr_text=%s, video_summary=%s, "
            "crawled_at=CURRENT_TIMESTAMP "
            "WHERE note_id=%s",
            (transcript, ocr_text, summary, note_id),
        )


def update_image_analysis(
    conn: Connection,
    note_id: str,
    ocr_text: str = "",
    summary: str = "",
    mermaid: str = "",
) -> None:
    """更新图片分析结果（OCR / AI 描述 / Mermaid），不影响其他字段。"""
    with conn:
        conn.execute(
            "UPDATE xhs_notes SET image_ocr_text=%s, image_summary=%s, image_mermaid=%s "
            "WHERE note_id=%s",
            (ocr_text, summary, mermaid, note_id),
        )


def upsert_user(conn: Connection, user: dict[str, Any]) -> None:
    user_id = user.get("user_id")
    if not user_id:
        return
    with conn:
        conn.execute(
            """INSERT INTO xhs_users (user_id, nickname, avatar, description,
                                  fans_count, follow_count, notes_count, location, raw_data)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
               ON CONFLICT(user_id) DO UPDATE SET
                   nickname = COALESCE(NULLIF(excluded.nickname, ''), xhs_users.nickname),
                   avatar = COALESCE(NULLIF(excluded.avatar, ''), xhs_users.avatar),
                   description = COALESCE(NULLIF(excluded.description, ''), xhs_users.description),
                   fans_count = CASE WHEN excluded.fans_count > 0 THEN excluded.fans_count ELSE xhs_users.fans_count END,
                   follow_count = CASE WHEN excluded.follow_count > 0 THEN excluded.follow_count ELSE xhs_users.follow_count END,
                   notes_count = CASE WHEN excluded.notes_count > 0 THEN excluded.notes_count ELSE xhs_users.notes_count END,
                   location = COALESCE(NULLIF(excluded.location, ''), xhs_users.location),
                   raw_data = CASE WHEN length(excluded.raw_data::text) > length(COALESCE(xhs_users.raw_data, '{}'::jsonb)::text) THEN excluded.raw_data ELSE xhs_users.raw_data END
            """,
            (
                user_id,
                user.get("nickname", ""),
                user.get("avatar", ""),
                user.get("description", ""),
                int(user.get("fans_count", 0) or 0),
                int(user.get("follow_count", 0) or 0),
                int(user.get("notes_count", 0) or 0),
                user.get("location", ""),
                json.dumps(user.get("raw", {}), ensure_ascii=False),
            ),
        )


def save_search_page(
    conn: Connection, keyword: str, page: int, note_ids: list[str]
) -> None:
    with conn:
        conn.execute(
            "INSERT INTO xhs_search_cache (keyword, page, note_ids) VALUES (%s, %s, %s) ON CONFLICT (keyword, page) DO UPDATE SET note_ids=EXCLUDED.note_ids, crawled_at=CURRENT_TIMESTAMP",
            (keyword, page, json.dumps(note_ids, ensure_ascii=False)),
        )


def update_crawl_state(
    conn: Connection,
    task_id: str,
    task_type: str,
    target_id: str,
    cursor: str,
    status: str,
    last_error: str = "",
) -> None:
    with conn:
        conn.execute(
            """INSERT INTO xhs_crawl_state
               (task_id, task_type, target_id, cursor, status, last_error, updated_at)
               VALUES (%s, %s, %s, %s, %s, %s, %s)
               ON CONFLICT (task_id) DO UPDATE SET task_type=EXCLUDED.task_type, target_id=EXCLUDED.target_id, cursor=EXCLUDED.cursor, status=EXCLUDED.status, last_error=EXCLUDED.last_error, updated_at=EXCLUDED.updated_at""",
            (task_id, task_type, target_id, cursor, status, last_error,
             datetime.now(timezone.utc).isoformat()),
        )


def get_crawl_state(conn: Connection, task_id: str) -> DictRow | None:
    cur = conn.execute("SELECT * FROM xhs_crawl_state WHERE task_id = %s", (task_id,))
    return cur.fetchone()


def get_note(conn: Connection, note_id: str) -> DictRow | None:
    cur = conn.execute("SELECT * FROM xhs_notes WHERE note_id = %s", (note_id,))
    return cur.fetchone()


def list_pending_corrections(conn: Connection, limit: int = 50) -> list[DictRow]:
    """待纠错视频笔记：有转录、有 OCR 参照、尚未标记纠错（video_summary 为空）。"""
    return conn.execute(
        "SELECT note_id, title, video_transcript, video_ocr_text FROM xhs_notes "
        "WHERE type='video' AND length(video_transcript) > 100 "
        "AND video_ocr_text != '' AND video_ocr_text != '[]' "
        "AND (video_summary = '' OR video_summary IS NULL) "
        "ORDER BY crawled_at DESC, note_id LIMIT %s",
        (limit,),
    ).fetchall()


def get_user(conn: Connection, user_id: str) -> DictRow | None:
    cur = conn.execute("SELECT * FROM xhs_users WHERE user_id = %s", (user_id,))
    return cur.fetchone()


def resolve_user_id(conn: Connection, name_or_id: str) -> str | None:
    """昵称 → user_id 查找。输入 user_id 则直接返回，输入昵称则查 DB。"""
    # 尝试作为 user_id 精确匹配
    row = conn.execute("SELECT user_id FROM xhs_users WHERE user_id = %s", (name_or_id,)).fetchone()
    if row:
        return row["user_id"]
    # 按昵称查找（精确优先，然后模糊）
    row = conn.execute("SELECT user_id FROM xhs_users WHERE nickname = %s", (name_or_id,)).fetchone()
    if row:
        return row["user_id"]
    row = conn.execute("SELECT user_id FROM xhs_users WHERE nickname LIKE %s", (f"%{name_or_id}%",)).fetchone()
    if row:
        return row["user_id"]
    return None


# ---------------------------------------------------------------------------
# Comments
# ---------------------------------------------------------------------------

def upsert_comment(conn: Connection, comment: dict[str, Any]) -> None:
    comment_id = comment.get("comment_id")
    if not comment_id:
        return
    with conn:
        conn.execute(
            """INSERT INTO xhs_comments (comment_id, note_id, parent_id, user_id, nickname,
                                     content, like_count, ip_location, pictures,
                                     target_comment_id, created_at, raw_data)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
               ON CONFLICT(comment_id) DO UPDATE SET
                   content = excluded.content,
                   like_count = CASE WHEN excluded.like_count > 0 THEN excluded.like_count ELSE xhs_comments.like_count END,
                   ip_location = COALESCE(NULLIF(excluded.ip_location, ''), xhs_comments.ip_location),
                   pictures = excluded.pictures,
                   raw_data = CASE WHEN length(excluded.raw_data::text) > length(COALESCE(xhs_comments.raw_data, '{}'::jsonb)::text) THEN excluded.raw_data ELSE xhs_comments.raw_data END
            """,
            (
                comment_id,
                comment.get("note_id"),
                comment.get("parent_id", ""),
                comment.get("user_id", ""),
                comment.get("nickname", ""),
                comment.get("content", ""),
                int(comment.get("like_count", 0) or 0),
                comment.get("ip_location", ""),
                json.dumps(_json_value(comment.get("pictures"), []), ensure_ascii=False),
                comment.get("target_comment_id", ""),
                comment.get("created_at", ""),
                json.dumps(comment.get("raw", {}), ensure_ascii=False),
            ),
        )


def iter_comments(conn: Connection, note_id: str) -> Iterable[DictRow]:
    """主评论按 like_count desc，子评论挂在主评论下面（按 created_at asc）"""
    main = conn.execute(
        "SELECT * FROM xhs_comments WHERE note_id = %s AND parent_id = '' "
        "ORDER BY like_count DESC, created_at ASC",
        (note_id,),
    ).fetchall()
    for m in main:
        yield m
        subs = conn.execute(
            "SELECT * FROM xhs_comments WHERE note_id = %s AND parent_id = %s "
            "ORDER BY created_at ASC",
            (note_id, m["comment_id"]),
        ).fetchall()
        yield from subs


def count_comments(conn: Connection, note_id: str) -> tuple[int, int]:
    """返回 (主评论数, 子评论数)"""
    main = conn.execute("SELECT COUNT(*) FROM xhs_comments WHERE note_id = %s AND parent_id = ''", (note_id,)).fetchone()[0]
    sub = conn.execute("SELECT COUNT(*) FROM xhs_comments WHERE note_id = %s AND parent_id != ''", (note_id,)).fetchone()[0]
    return main, sub


def iter_notes(
    conn: Connection, user_id: str | None = None
) -> Iterable[dict]:
    if user_id:
        cur = conn.execute(
            "SELECT * FROM xhs_notes WHERE user_id = %s ORDER BY published_at DESC",
            (user_id,),
        )
    else:
        cur = conn.execute("SELECT * FROM xhs_notes ORDER BY published_at DESC")
    for row in cur:
        yield dict(row)


def _note_context(conn: Connection, note_id: str, md_path: Path | None = None) -> dict:
    """提取渲染所需的公共数据，避免多个 _render_* 重复查库。"""
    note = get_note(conn, note_id)
    if not note:
        raise FileNotFoundError(f"note {note_id} not in DB")
    user = get_user(conn, note["user_id"]) if note["user_id"] else None
    raw = (note["raw_data"] or {})
    nickname = (user["nickname"] if user else "") or raw.get("user", {}).get("nickname", "")
    topics = (note["topics"] or [])
    topic_str = " ".join(f"#{t}" for t in topics) if topics else "—"

    # 计算从 MD 文件到 media 根目录的相对前缀
    media_rel_prefix = ".."
    if md_path is not None:
        try:
            depth = len(md_path.relative_to(OUTPUT_DIR).parts) - 1
            media_rel_prefix = "/".join([".."] * (depth + 1))
        except ValueError:
            pass

    media_dir = _find_media_dir(note_id, conn)
    local_images = sorted(media_dir.glob("img_*")) if media_dir.exists() else []
    local_video = next(iter(media_dir.glob("video.*")), None) if media_dir.exists() else None

    images = _extract_images(raw)
    if local_images:
        images = [str(Path(media_rel_prefix) / media_dir.relative_to(OUTPUT_DIR.parent) / p.name).replace("\\", "/") for p in local_images]
    video_url = _extract_video_url(raw)
    if local_video:
        video_url = str(Path(media_rel_prefix) / media_dir.relative_to(OUTPUT_DIR.parent) / local_video.name).replace("\\", "/")

    return {
        "note": note, "user": user, "raw": raw, "nickname": nickname,
        "topics": topics, "topic_str": topic_str, "media_dir": media_dir,
        "media_rel_prefix": media_rel_prefix,
        "images": images, "video_url": video_url,
        "local_video": local_video, "local_images": local_images,
    }


def _render_index(conn: Connection, ctx: dict) -> str:
    """渲染 index.md：元数据 + 正文 + 图片 + 视频基本信息。"""
    note = ctx["note"]
    lines = [
        f"# {note['title'] or '(无标题)'}",
        "",
        f"- **作者**: {ctx['nickname']} (@{note['user_id']})",
        f"- **发布时间**: {note['published_at'] or '—'}",
        f"- **IP属地**: {note['ip_location'] or '—'}",
        f"- **互动**: 赞 {note['liked_count']} | 藏 {note['collected_count']} | 评 {note['comment_count']} | 分享 {note['share_count']}",
        f"- **话题**: {ctx['topic_str']}",
        f"- **类型**: {note['type']}",
        f"- **笔记链接**: https://www.xiaohongshu.com/explore/{note['note_id']}",
        "",
        "## 正文",
        "",
        note["description"] or "(无正文)",
        "",
    ]
    if ctx["images"]:
        lines.append("## 图片")
        lines.append("")
        md_title = (note['title'] or '(无标题)').replace('[', '(').replace(']', ')')
        for i, url in enumerate(ctx["images"], 1):
            lines.append(f"![{md_title}·图{i}]({url})")
        lines.append("")

    # 视频基本信息（不含分析内容）
    if ctx["video_url"]:
        lines.append("## 视频")
        lines.append("")
        cover = note["cover_url"] or _extract_cover_url(ctx["raw"]) or ""
        if cover:
            local_cover = next(iter(ctx["media_dir"].glob("cover.*")), None) if ctx["media_dir"].exists() else None
            if local_cover:
                cover_ref = str(Path(ctx["media_rel_prefix"]) / ctx["media_dir"].relative_to(OUTPUT_DIR.parent) / local_cover.name).replace("\\", "/")
            else:
                cover_ref = cover
            lines.append(f"![封面]({cover_ref})")
            lines.append("")
        duration = note["video_duration"] or 0
        if duration:
            mins, secs = divmod(duration, 60)
            lines.append(f"- **时长**: {mins:02d}:{secs:02d}")
        lines.append(f"[视频链接]({ctx['video_url']})")
        if ctx["local_video"]:
            lines.append(f"- **本地文件**: `{ctx['local_video'].name}`")
        lines.append("")

    # 链接提示（指向其他子文件）
    links = []
    note_id = note["note_id"]
    transcript = note["video_transcript"] or ""
    ocr_text = note["video_ocr_text"] or ""
    summary = note["video_summary"] or ""
    image_summary = note["image_summary"] or ""
    image_ocr = note["image_ocr_text"] or ""
    image_mermaid = note["image_mermaid"] or ""
    main_n, _ = count_comments(conn, note_id)

    if summary or transcript or ocr_text:
        links.append("[视频分析](video.md)")
    if image_summary or image_ocr or image_mermaid:
        links.append("[图片分析](images.md)")
    if main_n:
        links.append(f"[评论](comments.md) ({main_n} 条)")

    if links:
        lines.append("---")
        lines.append("")
        for link in links:
            lines.append(f"- {link}")
        lines.append("")

    return "\n".join(lines)


def _render_video_section(ctx: dict) -> str | None:
    """渲染 video.md：摘要 + 转录 + OCR。无视频内容时返回 None。"""
    note = ctx["note"]
    transcript = note["video_transcript"] or ""
    ocr_text = note["video_ocr_text"] or ""
    summary = note["video_summary"] or ""

    if not (summary or transcript or ocr_text):
        return None

    title = note['title'] or '(无标题)'
    lines = [f"# {title} — 视频分析", ""]

    if summary:
        lines.append("## 摘要")
        lines.append("")
        lines.append(summary)
        lines.append("")
    elif transcript or ocr_text:
        lines.append("## 内容提取结果")
        lines.append("")
        lines.append(f"- 转录: {len(transcript)} 字")
        try:
            ocr_count = len(json.loads(ocr_text))
        except (json.JSONDecodeError, TypeError):
            ocr_count = 0
        lines.append(f"- 画面文字: {ocr_count} 帧")
        lines.append("")

    if transcript:
        lines.append("## 语音转录")
        lines.append("")
        lines.append(transcript)
        lines.append("")

    ocr_rendered = _render_video_ocr(ocr_text)
    if ocr_rendered:
        lines.append("## 画面文字")
        lines.append("")
        lines.append(ocr_rendered)
        lines.append("")

    return "\n".join(lines)


def _render_images_section(ctx: dict) -> str | None:
    """渲染 images.md：AI 描述 + OCR + Mermaid。无图片分析时返回 None。"""
    note = ctx["note"]
    image_summary = note["image_summary"] or ""
    image_ocr = note["image_ocr_text"] or ""
    image_mermaid = note["image_mermaid"] or ""

    if not (image_summary or image_ocr or image_mermaid):
        return None

    title = note['title'] or '(无标题)'
    lines = [f"# {title} — 图片分析", ""]

    if image_summary:
        lines.append("## AI 描述")
        lines.append("")
        lines.append(image_summary)
        lines.append("")
    if image_ocr:
        lines.append("## 图片文字")
        lines.append("")
        lines.append(image_ocr)
        lines.append("")
    if image_mermaid:
        lines.append("## 路线图 / 流程图")
        lines.append("")
        lines.append("```mermaid")
        lines.append(image_mermaid)
        lines.append("```")
        lines.append("")

    return "\n".join(lines)


def _render_comments_section(conn: Connection, note_id: str) -> str | None:
    """渲染 comments.md：全部评论。无评论时返回 None。"""
    main_n, sub_n = count_comments(conn, note_id)
    if not main_n:
        return None

    note = get_note(conn, note_id)
    title = (note["title"] or "(无标题)") if note else note_id
    lines = [f"# {title} — 评论", ""]
    lines.append(f"共 {main_n} 条主评论 + {sub_n} 条回复")
    lines.append("")

    for c in iter_comments(conn, note_id):
        indent = "" if c["parent_id"] == "" else "  "
        head = f"**{c['nickname']}** ({c['ip_location'] or '?'}, 赞 {c['like_count']})"
        lines.append(f"{indent}- {head}: {c['content']}")
    lines.append("")

    return "\n".join(lines)


def write_markdown_files(conn: Connection, note_id: str) -> list[Path]:
    """为单篇笔记生成子文件目录：index.md + video.md + images.md + comments.md。"""
    # 博主子目录
    note = get_note(conn, note_id)
    author_dir = OUTPUT_DIR
    if note and note["user_id"]:
        user = get_user(conn, note["user_id"])
        nickname = (user["nickname"] if user else "") or ""
        if nickname:
            author_dir = OUTPUT_DIR / sanitize_filename(nickname, 30)

    # 笔记子目录
    slug = _human_filename(conn, note_id, "")
    note_dir = author_dir / slug
    note_dir.mkdir(parents=True, exist_ok=True)

    # 公共数据
    index_path = note_dir / "index.md"
    ctx = _note_context(conn, note_id, md_path=index_path)

    files = []

    # index.md（必有）
    index_content = _render_index(conn, ctx)
    _atomic_write(index_path, index_content)
    files.append(index_path)

    # video.md（有视频分析时生成）
    video_content = _render_video_section(ctx)
    if video_content:
        path = note_dir / "video.md"
        _atomic_write(path, video_content)
        files.append(path)

    # images.md（有图片分析时生成）
    images_content = _render_images_section(ctx)
    if images_content:
        path = note_dir / "images.md"
        _atomic_write(path, images_content)
        files.append(path)

    # comments.md（有评论时生成）
    comments_content = _render_comments_section(conn, note_id)
    if comments_content:
        path = note_dir / "comments.md"
        _atomic_write(path, comments_content)
        files.append(path)

    return files


def _render_video_ocr(ocr_text: str) -> str:
    """智能渲染视频 OCR：JSON 解析 → 去重 → 噪声过滤。旧数据纯文本直接返回。"""
    if not ocr_text:
        return ""
    try:
        items = json.loads(ocr_text)
    except (json.JSONDecodeError, TypeError):
        return ocr_text  # 旧数据兼容

    seen = set()
    lines = []
    for item in items:
        text = item.get("text", "").strip()
        if not text or len(text) < 2 or text in seen:
            continue
        seen.add(text)
        time_label = item.get("time", "")
        if time_label:
            lines.append(f"- [{time_label}] {text}")
        else:
            lines.append(f"- {text}")
    return "\n".join(lines) if lines else ""


def write_markdown(conn: Connection, note_id: str) -> Path:
    """导出笔记为 MD 子文件目录，返回 index.md 路径。"""
    files = write_markdown_files(conn, note_id)
    return files[0]  # index.md


def _human_filename(conn: Connection, note_id: str, ext: str) -> str:
    """生成人类可读的文件名：{标题}_{note_id前8位}.ext

    无标题时退化到 {note_id}.ext。
    """
    row = get_note(conn, note_id)
    if not row:
        return f"{note_id}{ext}"
    title = row["title"] or ""
    if not title:
        desc = row["description"] or ""
        title = desc[:50] if desc else ""
    # 清理标题：去掉文件名不合法字符
    title = re.sub(r'[<>:"/\\|?*\n\r\t]', '', title)
    title = title.replace(' ', '_')
    title = title[:70]  # 截断
    if not title:
        return f"{note_id}{ext}"
    return f"{title}_{note_id[:8]}{ext}"


def render_update_summary(conn: Connection, note_id: str) -> str:
    """生成笔记当前状态摘要，用于更新通知。"""
    row = get_note(conn, note_id)
    if not row:
        return ""
    parts = []
    title = row["title"] or "(无标题)"

    # 基本信息
    type_label = "视频" if row["type"] == "video" else "图文"
    parts.append(f"《{title}》({type_label})")

    # 媒体状态
    media_dir = _find_media_dir(note_id, conn)
    local_imgs = len(list(media_dir.glob("img_*"))) if media_dir.exists() else 0
    local_video = next(iter(media_dir.glob("video.*")), None) is not None if media_dir.exists() else False
    media_items = []
    if local_imgs:
        media_items.append(f"{local_imgs} 张图片")
    if local_video:
        media_items.append("1 个视频")
    if media_items:
        parts.append("已下载: " + "、".join(media_items))

    # 分析状态
    analysis_items = []
    if row["video_transcript"]:
        analysis_items.append(f"语音转录 {len(row['video_transcript'])} 字")
    if row["video_ocr_text"]:
        analysis_items.append(f"画面文字 {len(row['video_ocr_text'])} 字")
    if row["video_summary"]:
        analysis_items.append("AI 摘要")
    if analysis_items:
        parts.append("视频分析: " + "、".join(analysis_items))

    # 图片分析状态
    image_items = []
    if row["image_ocr_text"]:
        image_items.append(f"图片OCR {len(row['image_ocr_text'])} 字")
    if row["image_summary"]:
        image_items.append("图片AI描述")
    if row["image_mermaid"]:
        image_items.append("路线图/流程图")
    if image_items:
        parts.append("图片分析: " + "、".join(image_items))

    # 评论
    main_n, sub_n = count_comments(conn, note_id)
    if main_n:
        parts.append(f"{main_n} 条评论")

    return " | ".join(parts)


CSV_HEADERS = [
    "序号", "标题", "作者", "发布时间", "类型",
    "点赞", "收藏", "评论", "分享", "IP属地",
    "话题", "笔记链接", "媒体", "正文摘要",
]


def write_csv(conn: Connection, path: Path | None = None, user_id: str | None = None) -> list[Path]:
    """导出 CSV。按博主分文件，放到 output/<博主名>/ 目录。返回生成的文件列表。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 按 user_id 分组
    if user_id:
        groups = {user_id: list(iter_notes_by_user(conn, user_id))}
    else:
        groups: dict[str, list] = {}
        for note in iter_notes(conn):
            uid = note["user_id"] or "_unknown"
            groups.setdefault(uid, []).append(note)

    files = []
    for uid, notes in groups.items():
        # 确定博主目录
        user = get_user(conn, uid) if uid != "_unknown" else None
        nickname = (user["nickname"] if user else "") or uid
        safe_author = sanitize_filename(nickname, 30)
        author_dir = OUTPUT_DIR / safe_author
        author_dir.mkdir(parents=True, exist_ok=True)

        csv_path = path or (author_dir / f"{safe_author}_笔记列表.csv")
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, quoting=csv.QUOTE_ALL)
            writer.writerow(CSV_HEADERS)
            for i, note in enumerate(notes, 1):
                note_user = get_user(conn, note["user_id"]) if note["user_id"] else None
                topics = (note["topics"] or [])
                desc = note["description"] or ""

                # 媒体摘要
                csv_media_dir = _find_media_dir(note["note_id"], conn)
                media_parts = []
                if csv_media_dir.exists():
                    img_count = len(list(csv_media_dir.glob("img_*")))
                    has_video = any(csv_media_dir.glob("video.*"))
                    if has_video:
                        media_parts.append("视频")
                    if img_count:
                        media_parts.append(f"{img_count}张图")
                media_str = "+".join(media_parts) if media_parts else "—"

                # 正文摘要：截断 + 换行替换
                summary = desc[:200].replace("\n", " ").replace("\r", " ").strip()

                writer.writerow([
                    i,
                    note["title"] or "",
                    (note_user["nickname"] if note_user else "") or "",
                    note["published_at"] or "",
                    "视频" if note["type"] == "video" else "图文",
                    note["liked_count"],
                    note["collected_count"],
                    note["comment_count"],
                    note["share_count"],
                    note["ip_location"] or "",
                    " ".join(f"#{t}" for t in topics),
                    f"https://www.xiaohongshu.com/explore/{note['note_id']}",
                    media_str,
                    summary,
                ])
        files.append(csv_path)
        path = None  # 后续分组自动生成路径

    return files


def iter_notes_by_user(conn: Connection, user_id: str):
    """按 user_id 迭代笔记。"""
    cur = conn.execute(
        "SELECT * FROM xhs_notes WHERE user_id = %s ORDER BY published_at DESC",
        (user_id,),
    )
    for row in cur:
        yield dict(row)


def write_json(conn: Connection, path: Path | None = None, user_id: str | None = None) -> Path:
    """导出笔记为 JSON 文件。user_id 过滤指定博主。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    date_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    if user_id:
        user = get_user(conn, user_id)
        nickname = (user["nickname"] if user else "") or user_id
        safe_author = sanitize_filename(nickname, 30)
        author_dir = OUTPUT_DIR / safe_author
        author_dir.mkdir(parents=True, exist_ok=True)
        path = path or (author_dir / f"{safe_author}_笔记.json")
    else:
        path = path or (OUTPUT_DIR / f"xhs_export_{date_str}.json")
    notes = []
    source = iter_notes_by_user(conn, user_id) if user_id else iter_notes(conn)
    for note in source:
        d = dict(note)
        d["topics"] = d.get("topics") or []
        d["raw_data"] = d.get("raw_data") or {}
        notes.append(d)
    _atomic_write(path, json.dumps(notes, ensure_ascii=False, indent=2, default=str))
    return path


def write_xlsx(conn: Connection, path: Path | None = None, user_id: str | None = None) -> Path:
    """导出笔记为 XLSX 文件（多 sheet）。user_id 过滤指定博主。需要 openpyxl。"""
    try:
        import openpyxl  # type: ignore
    except ImportError:
        print("[ERR] 导出 xlsx 需要 openpyxl。pip install openpyxl", file=sys.stderr)
        raise
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if user_id:
        user = get_user(conn, user_id)
        nickname = (user["nickname"] if user else "") or user_id
        safe_author = sanitize_filename(nickname, 30)
        author_dir = OUTPUT_DIR / safe_author
        author_dir.mkdir(parents=True, exist_ok=True)
        date_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        path = path or (author_dir / f"{safe_author}_笔记.xlsx")
    else:
        date_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        path = path or (OUTPUT_DIR / f"xhs_export_{date_str}.xlsx")
    wb = openpyxl.Workbook()

    # Sheet 1: Notes
    ws_notes = wb.active
    ws_notes.title = "Notes"
    note_cols = [
        "note_id", "user_id", "title", "description", "type",
        "liked_count", "collected_count", "comment_count", "share_count",
        "ip_location", "published_at", "xsec_token",
        "video_url", "cover_url", "video_duration",
        "video_summary", "image_ocr_text", "image_summary",
        "crawled_at",
    ]
    ws_notes.append(note_cols)
    for note in (iter_notes_by_user(conn, user_id) if user_id else iter_notes(conn)):
        ws_notes.append([value.isoformat() if isinstance(value, datetime) else value for value in (note.get(c, "") for c in note_cols)])

    # Sheet 2: Users
    ws_users = wb.create_sheet("Users")
    user_cols = ["user_id", "nickname", "avatar", "fans_count", "notes_count", "desc"]
    ws_users.append(user_cols)
    for row in conn.execute("SELECT * FROM xhs_users ORDER BY fans_count DESC"):
        d = dict(row)
        ws_users.append([d.get(c, "") for c in user_cols])

    # Sheet 3: Comments
    ws_comments = wb.create_sheet("Comments")
    comment_cols = ["comment_id", "note_id", "content", "nickname", "like_count", "created_at"]
    ws_comments.append(comment_cols)
    for row in conn.execute("SELECT * FROM xhs_comments ORDER BY created_at DESC"):
        d = dict(row)
        ws_comments.append([d.get(c, "") for c in comment_cols])

    wb.save(path)
    return path


def _extract_images(raw: dict) -> list[str]:
    images: list[str] = []
    for img in raw.get("image_list", []) or []:
        url = img.get("url_default") or img.get("url") or ""
        info_list = img.get("info_list") or []
        for info in info_list:
            if info.get("image_scene") in ("WB_DFT", "WB_PRV"):
                url = info.get("url") or url
                break
        if url:
            images.append(url)
    return images


def _extract_video_url(raw: dict) -> str | None:
    video = raw.get("video") or {}
    media = video.get("media") or {}
    stream = media.get("stream") or {}
    for codec_key in ("h264", "h265", "av1"):
        items = stream.get(codec_key) or []
        if items:
            return items[0].get("master_url") or items[0].get("backup_urls", [None])[0]
    return None


def _extract_cover_url(raw: dict) -> str:
    """从 raw_data 提取视频封面图 URL。"""
    video = raw.get("video") or {}
    # 尝试多种路径
    for path in [
        lambda: video.get("cover", ""),
        lambda: (video.get("image_list") or [{}])[0].get("url_default", "") if video.get("image_list") else "",
        lambda: (raw.get("image_list") or [{}])[0].get("url_default", "") if raw.get("image_list") else "",
    ]:
        try:
            url = path()
            if url:
                return url
        except (IndexError, TypeError):
            continue
    return ""


def _extract_video_duration(raw: dict) -> int:
    """从 raw_data 提取视频时长（秒）。"""
    video = raw.get("video") or {}
    dur = video.get("duration")
    if dur:
        return int(dur) // 1000 if int(dur) > 1000 else int(dur)
    return 0


def _find_media_dir(note_id: str, conn: Connection) -> Path:
    """查找笔记的本地媒体目录（兼容新旧两种结构）。

    新结构: media/<博主名>/<笔记标题>/
    旧结构: media/<note_id>/

    当新旧目录都存在时，优先返回有实际内容的那个；
    如果都有内容，优先新路径（后续下载会写入新路径）。
    """
    new_dir = note_media_dir(note_id, conn)
    new_has_content = new_dir.exists() and any(new_dir.iterdir())
    legacy_dir = MEDIA_DIR / note_id
    legacy_has_content = legacy_dir.exists() and any(legacy_dir.iterdir())

    if new_has_content:
        return new_dir
    if legacy_has_content:
        return legacy_dir
    # 都不存在或都为空，返回新路径（作为目标目录）
    return new_dir
